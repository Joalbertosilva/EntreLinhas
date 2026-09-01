#!/usr/bin/env python3
"""Gera EntreLinhas-Stack-e-Ferramentas.docx a partir do conteúdo de stack."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "EntreLinhas-Stack-e-Ferramentas.docx"

# Paleta EntreLinhas
BRAND = RGBColor(0x1A, 0x5C, 0x6E)       # teal escuro
BRAND_LIGHT = RGBColor(0x2A, 0x8F, 0xA8)
ACCENT = RGBColor(0xE8, 0x6C, 0x3A)      # laranja
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = "F3F4F6"
TABLE_HEADER = "1A5C6E"
TABLE_ALT = "F0F9FB"
QUOTE_BG = "E8F4F8"


def set_cell_shading(cell, fill_hex: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, fill_hex: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    p_pr.append(shd)


def style_run(run, *, bold=False, italic=False, size=11, color=TEXT, font="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font


def add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A5C6E")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("EntreLinhas")
    style_run(r, bold=True, size=36, color=BRAND, font="Georgia")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Stack, Ferramentas e Arquitetura\nWeb + App Mobile + Backend")
    style_run(r, size=20, color=BRAND_LIGHT, font="Georgia")

    doc.add_paragraph()
    doc.add_paragraph()

    meta_lines = [
        ("Projeto", "Plataforma de estímulo à leitura, reflexão e produção textual"),
        ("Contexto", "Trabalho de Conclusão de Curso (TCC) — CESMAC"),
        ("Repositório", "Monorepo tcc-sistema"),
        ("Atualização", "Setembro / 2026"),
    ]
    for label, value in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}: ")
        style_run(r1, bold=True, size=11, color=MUTED)
        r2 = p.add_run(value)
        style_run(r2, size=11, color=TEXT)

    doc.add_paragraph()
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tag.add_run("Documento técnico de referência")
    style_run(r, italic=True, size=10, color=MUTED)

    doc.add_page_break()


def add_toc(doc: Document) -> None:
    h = doc.add_heading("Sumário", level=1)
    for run in h.runs:
        style_run(run, bold=True, size=16, color=BRAND)

    sections = [
        "1. Visão geral (web + mobile + backend)",
        "2. Linguagem principal: TypeScript",
        "3. Monorepo com pnpm workspaces",
        "4. Frontend web (apps/web)",
        "5. App mobile (apps/mobile) — planejado",
        "6. Supabase — backend como serviço",
        "7. Pacotes compartilhados",
        "8. Como uma funcionalidade é construída",
        "9. Ferramentas de desenvolvimento e qualidade",
        "10. Estado atual do desenvolvimento",
        "11. Por que essas escolhas? (síntese para o TCC)",
        "12. Referências",
    ]
    for item in sections:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            style_run(run, size=11, color=TEXT)

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: BRAND, 2: BRAND_LIGHT, 3: TEXT}
    for run in h.runs:
        style_run(run, bold=True, size=sizes.get(level, 11), color=colors.get(level, TEXT))
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for run in p.runs:
        style_run(run, size=11, color=TEXT)


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(p, QUOTE_BG)
    r = p.add_run(text)
    style_run(r, italic=True, size=10, color=MUTED)


def add_code_block(doc: Document, text: str) -> None:
    for line in text.strip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_paragraph_shading(p, CODE_BG)
        r = p.add_run(line if line else " ")
        style_run(r, size=9, color=TEXT, font="Consolas")


def add_diagram_box(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(0)
        set_paragraph_shading(p, CODE_BG)
        r = p.add_run(line)
        style_run(r, size=9, color=BRAND, font="Consolas")
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_shading(cell, TABLE_HEADER)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        style_run(r, bold=True, size=10, color=WHITE)

    for row_idx, row in enumerate(rows):
        cells = table.rows[row_idx + 1].cells
        for col_idx, value in enumerate(row):
            if row_idx % 2 == 1:
                set_cell_shading(cells[col_idx], TABLE_ALT)
            p = cells[col_idx].paragraphs[0]
            r = p.add_run(value)
            style_run(r, size=10, color=TEXT)

    doc.add_paragraph()


def add_status_table(doc: Document, rows: list[list[str]]) -> None:
    headers = ["Módulo", "Status"]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, TABLE_HEADER)
        r = cell.paragraphs[0].add_run(header)
        style_run(r, bold=True, size=10, color=WHITE)

    for row_idx, (module, status) in enumerate(rows):
        cells = table.rows[row_idx + 1].cells
        if row_idx % 2 == 1:
            set_cell_shading(cells[0], TABLE_ALT)
            set_cell_shading(cells[1], TABLE_ALT)
        r0 = cells[0].paragraphs[0].add_run(module)
        style_run(r0, size=10, color=TEXT)
        r1 = cells[1].paragraphs[0].add_run(status)
        if status.startswith("✅"):
            style_run(r1, size=10, color=RGBColor(0x16, 0x65, 0x34))
        else:
            style_run(r1, size=10, color=RGBColor(0xB4, 0x53, 0x09))
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "1. Visão geral")
    add_body(
        doc,
        "O EntreLinhas é composto por três camadas no mesmo monorepo: website/gerenciador (apps/web), "
        "app mobile para alunos (apps/mobile — Fase 3) e backend Supabase compartilhado. "
        "Professores e administradores usam só a web; alunos poderão usar web ou app.",
    )
    add_table(
        doc,
        ["Camada", "Pasta", "Público", "Status"],
        [
            ["Website + gerenciador", "apps/web", "Alunos (/app) + staff (/admin)", "✅ em desenvolvimento"],
            ["App mobile", "apps/mobile", "Alunos (Android/iOS)", "⏳ Fase 3"],
            ["Backend", "supabase/", "Todas as plataformas", "✅ pronto"],
        ],
    )
    add_heading(doc, "Web vs mobile — quem usa o quê", level=2)
    add_table(
        doc,
        ["Funcionalidade", "Web /app", "Web /admin", "App mobile"],
        [
            ["Login", "✅", "✅", "⏳"],
            ["Vitrines e leitura", "✅", "—", "⏳"],
            ["Reflexão e comentários", "✅", "—", "⏳"],
            ["CRUD de conteúdos", "—", "✅", "—"],
            ["Gestão de usuários", "—", "✅", "—"],
        ],
    )
    add_body(doc, "Arquitetura (web + mobile + backend):")
    add_diagram_box(
        doc,
        [
            "┌─────────────────────┐     ┌─────────────────────┐",
            "│   apps/mobile       │     │      apps/web        │",
            "│   React Native      │     │   React + Vite       │",
            "│   Expo · iOS/Android│     │   /app + /admin      │",
            "└──────────┬──────────┘     └──────────┬───────────┘",
            "           │                           │",
            "           │  Supabase JS · TanStack Query · types/schemas",
            "           └─────────────┬─────────────┘",
            "                         │ HTTPS",
            "                         ▼",
            "┌─────────────────────────────────────────────────────────────┐",
            "│  Supabase — PostgreSQL · Auth · Storage · Edge Functions · RLS│",
            "└─────────────────────────────────────────────────────────────┘",
        ],
    )
    add_quote(doc, "Ordem de entrega: Backend → Web → Mobile (seção 10).")
    add_horizontal_rule(doc)

    add_heading(doc, "2. Linguagem principal: TypeScript")
    add_body(doc, "Frontend web, app mobile e Edge Functions são TypeScript.")
    add_table(
        doc,
        ["Onde", "Linguagem"],
        [
            ["Frontend web", "TypeScript + TSX (React)"],
            ["App mobile", "TypeScript + TSX (React Native)"],
            ["Pacotes compartilhados", "TypeScript"],
            ["Edge Functions (Supabase)", "TypeScript (runtime Deno)"],
            ["Banco de dados", "SQL (migrations versionadas)"],
            ["Scripts auxiliares", "JavaScript (Node)"],
        ],
    )
    add_quote(doc, "Tipos e schemas compartilhados entre web, mobile e backend.")
    add_horizontal_rule(doc)

    add_heading(doc, "3. Monorepo com pnpm workspaces")
    add_body(doc, "O projeto vive em um único repositório com vários pacotes interligados:")
    add_code_block(
        doc,
        """tcc-sistema/
├── apps/web/              # Website + gerenciador
├── apps/mobile/           # App aluno (Expo) — Fase 3
├── packages/types/        # Tipos compartilhados
├── packages/schemas/      # Validação Zod
├── supabase/migrations/   # SQL versionado
└── supabase/functions/    # Edge Functions""",
    )
    add_table(
        doc,
        ["Ferramenta", "Função"],
        [
            ["pnpm", "Gerenciador de pacotes — monorepo nativo, mais rápido que npm"],
            ["pnpm workspaces", "Pacotes @tcc-sistema/types e @tcc-sistema/schemas importados pelo web"],
            ["Node.js ≥ 20", "Runtime para desenvolvimento e build"],
        ],
    )
    add_heading(doc, "Comandos frequentes", level=2)
    add_code_block(
        doc,
        """pnpm web:dev          # Frontend em http://localhost:5173
pnpm web:build        # Build de produção
pnpm db:start         # Supabase local (Docker)
pnpm db:migrate       # Aplica migrations SQL pendentes
pnpm functions:serve  # Edge Functions locais""",
    )
    add_horizontal_rule(doc)

    add_heading(doc, "4. Frontend web (apps/web)")

    add_heading(doc, "4.1 Framework e build", level=2)
    add_table(
        doc,
        ["Tecnologia", "Versão", "Papel"],
        [
            ["React", "19.x", "Interface em componentes"],
            ["Vite", "8.x", "Dev server rápido + bundler de produção"],
            ["TypeScript", "6.x", "Tipagem estática"],
        ],
    )

    add_heading(doc, "4.2 Roteamento — TanStack Router", level=2)
    add_body(doc, "Rotas baseadas em arquivos em apps/web/src/routes/. Guards redirecionam aluno → /app e staff → /admin.")
    add_table(
        doc,
        ["Rota", "Descrição"],
        [
            ["/login", "Entrada no sistema"],
            ["/app", "Home do aluno"],
            ["/app/livros, /app/cronicas…", "Vitrines por tipo de conteúdo"],
            ["/app/conteudos/$id", "Detalhe — leitura, reflexão, comentários"],
            ["/admin", "Dashboard staff"],
            ["/admin/conteudos", "CRUD de conteúdos"],
            ["/admin/usuarios", "Gestão de usuários"],
        ],
    )

    add_heading(doc, "4.3 Dados no cliente — TanStack Query", level=2)
    add_body(
        doc,
        "Toda leitura/escrita ao Supabase passa por hooks com TanStack Query: cache automático, "
        "mutations para salvar reflexão/curtir/leitura, e invalidação após edições no admin. "
        "Exemplos: useConteudos.ts, useConteudoDetail.ts, useConteudoEngagement.ts.",
    )

    add_heading(doc, "4.4 Formulários e validação", level=2)
    add_table(
        doc,
        ["Biblioteca", "Uso"],
        [
            ["React Hook Form", "Formulários performáticos no admin"],
            ["Zod", "Regras de validação em @tcc-sistema/schemas"],
            ["@hookform/resolvers", "Integração Zod + React Hook Form"],
        ],
    )

    add_heading(doc, "4.5 Estilo visual", level=2)
    add_table(
        doc,
        ["Tecnologia", "Uso"],
        [
            ["Tailwind CSS 4", "Utility-first CSS"],
            ["CSS customizado", "Tokens de marca, leitor paginado, animações"],
            ["Lucide React", "Ícones"],
            ["Sonner", "Toasts de feedback"],
        ],
    )
    add_body(doc, "Componentes UI em apps/web/src/components/ui/ — padrão inspirado em shadcn/ui, adaptado ao design EntreLinhas.")

    add_heading(doc, "4.6 Organização do código", level=2)
    add_code_block(
        doc,
        """apps/web/src/
├── routes/        # Páginas (TanStack Router)
├── components/    # Layouts e UI genérica
├── features/      # Lógica por domínio
│   ├── app/       # Plataforma aluno
│   ├── conteudos/ # Admin de conteúdos
│   └── auth/      # Login e sessão
└── lib/           # Supabase client, utils""",
    )
    add_horizontal_rule(doc)

    add_heading(doc, "5. App mobile (apps/mobile) — planejado")
    add_quote(
        doc,
        "Fase 3 do projeto. Será iniciado após a plataforma web do aluno estar funcional. "
        "O backend Supabase já atende web e mobile — não haverá outro servidor.",
    )

    add_heading(doc, "5.1 Público e escopo", level=2)
    add_table(
        doc,
        ["Item", "Definição"],
        [
            ["Público", "Alunos (Android e iOS)"],
            ["Fora do app", "Gerenciador — permanece só na web (/admin)"],
            ["Paridade", "Mesmo fluxo da área /app da web"],
            ["Requisito", "RNF09 — consistência visual e funcional app/web"],
        ],
    )

    add_heading(doc, "5.2 Stack mobile", level=2)
    add_table(
        doc,
        ["Tecnologia", "Papel"],
        [
            ["React Native + Expo", "App nativo iOS/Android; Expo Go para testes"],
            ["Expo Router", "Rotas baseadas em arquivos"],
            ["NativeWind", "Estilo utility-first (Tailwind no mobile)"],
            ["TanStack Query", "Cache e requisições ao Supabase"],
            ["React Hook Form + Zod", "Formulários — @tcc-sistema/schemas"],
            ["Lucide React Native", "Ícones (mesma família da web)"],
            ["expo-secure-store", "Sessão JWT segura no dispositivo"],
            ["expo-image", "Capas otimizadas"],
        ],
    )

    add_heading(doc, "5.3 Rotas previstas (Expo Router)", level=2)
    add_code_block(
        doc,
        """apps/mobile/app/
├── (auth)/login.tsx
└── (aluno)/
    ├── index.tsx              # Home
    ├── conteudos/[id].tsx     # Detalhe — leitura, reflexão
    ├── leitura/
    ├── producoes/
    └── perfil.tsx""",
    )

    add_heading(doc, "5.4 Conexão com o backend", level=2)
    add_body(
        doc,
        "O app usa o mesmo Supabase da web: @supabase/supabase-js → PostgREST → PostgreSQL + RLS. "
        "Hooks espelham apps/web/src/features/app/. Variáveis: EXPO_PUBLIC_SUPABASE_URL e EXPO_PUBLIC_SUPABASE_ANON_KEY.",
    )

    add_heading(doc, "5.5 Reutilização vs reescrita", level=2)
    add_table(
        doc,
        ["Camada", "Reutilização"],
        [
            ["packages/types e schemas", "✅ 100% compartilhado"],
            ["Hooks Supabase (queries/mutations)", "🔄 Mesma lógica; UI em RN"],
            ["Componentes de UI", "❌ Reescritos em React Native"],
            ["Design system", "🔄 Mesmos tokens de marca"],
        ],
    )

    add_heading(doc, "5.6 Desenvolvimento e distribuição", level=2)
    add_code_block(
        doc,
        """pnpm db:start                    # Terminal 1
pnpm --filter mobile start       # Terminal 2 — Expo Dev Tools""",
    )
    add_table(
        doc,
        ["Etapa", "Ferramenta"],
        [
            ["Desenvolvimento", "Expo Go + emuladores"],
            ["Build Android", "EAS Build → Google Play"],
            ["Build iOS", "EAS Build → TestFlight / App Store"],
            ["Backend produção", "Mesmo Supabase da web"],
        ],
    )
    add_horizontal_rule(doc)

    add_heading(doc, "6. Supabase — backend como serviço")

    add_heading(doc, "6.1 O que é o Supabase", level=2)
    add_body(doc, "Plataforma open-source baseada em PostgreSQL que entrega:")
    for item in [
        "Banco PostgreSQL — tabelas relacionais",
        "API REST automática (PostgREST) — client JS fala direto com tabelas permitidas",
        "Auth — login, JWT, sessão",
        "Storage — arquivos (capas de livros)",
        "Edge Functions — operações privilegiadas em TypeScript/Deno",
        "RLS — políticas SQL que filtram linhas por usuário/perfil",
    ]:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            style_run(run, size=10, color=TEXT)

    add_heading(doc, "6.2 Desenvolvimento local", level=2)
    add_code_block(doc, "pnpm db:start    # Docker: Postgres + Auth + Storage + Studio\npnpm db:migrate  # Aplica supabase/migrations/")

    add_heading(doc, "6.3 Migrations (SQL versionado)", level=2)
    add_table(
        doc,
        ["Migration", "Conteúdo"],
        [
            ["20260825120000_enums_and_helpers.sql", "Enums e funções auxiliares"],
            ["20260825120100_profiles.sql", "Perfil do usuário"],
            ["20260825120200_conteudos.sql", "Conteúdos, temas, materiais"],
            ["20260825120300_interacoes_leituras.sql", "Interações e progresso de leitura"],
            ["20260831180000_curtidas_interacoes_publicas.sql", "Curtidas e comentários públicos"],
        ],
    )

    add_heading(doc, "6.4 Principais tabelas", level=2)
    add_table(
        doc,
        ["Tabela", "Finalidade"],
        [
            ["profiles", "Nome, nome_usuario, perfil (aluno/professor/admin)"],
            ["conteudos", "Livros, crônicas, poemas — texto, capa, metadados"],
            ["temas", "Reflexão orientada (frase, orientação, pergunta)"],
            ["materiais_complementares", "Links, vídeos, áudios"],
            ["interacoes", "Reflexões privadas e comentários públicos"],
            ["leituras", "Status em_andamento / concluido"],
            ["conteudo_curtidas", "Curtidas para seção Destaques"],
            ["audit_logs", "Auditoria de ações administrativas"],
            ["password_reset_requests", "Recuperação de senha"],
        ],
    )

    add_heading(doc, "6.5 Row Level Security (RLS)", level=2)
    add_body(doc, "Regras no PostgreSQL definem o que cada usuário pode ler/escrever:")
    for item in [
        "Aluno lê conteúdos ativos; vê só suas reflexões; vê comentários públicos de outros",
        "Staff gerencia conteúdos, usuários e acompanha reflexões dos alunos",
        "Storage — upload de capas só staff; leitura conforme política do bucket",
    ]:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            style_run(run, size=10, color=TEXT)

    add_heading(doc, "6.6 Autenticação", level=2)
    add_table(
        doc,
        ["Aspecto", "Implementação"],
        [
            ["Login", "Nome de usuário + senha (mapeado para Supabase Auth)"],
            ["Sessão", "JWT guardado pelo client Supabase"],
            ["Cadastro", "Somente admin via Edge Function create-user"],
            ["Senha", "Recuperação com request-password-reset e admin-reset-password"],
        ],
    )

    add_heading(doc, "6.7 Storage e Edge Functions", level=2)
    add_body(doc, "Bucket de capas: upload no gerenciador, URL em conteudos.capa_url.")
    add_table(
        doc,
        ["Edge Function", "Quando usar"],
        [
            ["create-user", "Admin cadastra aluno/professor"],
            ["request-password-reset", "Aluno solicita redefinição"],
            ["admin-reset-password", "Staff aprova e define senha temporária"],
        ],
    )
    add_horizontal_rule(doc)

    add_heading(doc, "7. Pacotes compartilhados")
    add_heading(doc, "@tcc-sistema/types", level=2)
    add_body(doc, "Tipos TypeScript espelhando o banco: Conteudo, Tema, Interacao, Profile; enums TipoConteudo, StatusLeitura, TipoInteracao.")
    add_heading(doc, "@tcc-sistema/schemas", level=2)
    add_body(doc, "Schemas Zod: conteudoFormSchema, temaFormSchema, loginSchema. Testes com Vitest (pnpm test:schemas).")
    add_horizontal_rule(doc)

    add_heading(doc, "8. Como uma funcionalidade é construída (exemplo)")
    add_body(doc, "Fluxo: cadastrar reflexão de um livro e o aluno responder.")
    steps = [
        "Admin preenche ContentFormDialog.tsx — React Hook Form + Zod",
        "Save grava conteudos e sincroniza temas (conteudoReflexao.ts)",
        "Aluno abre /app/conteudos/$id — useConteudoDetail busca dados",
        "Leitor paginado (LivroDetailPage, BookPaginatedReader) — uma seção por página",
        "Reflexão gravada em interacoes (tipo reflexao_orientada); RLS garante privacidade",
        "Comentário na mesma tabela (tipo comentario_livre), visível entre alunos",
    ]
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(f"{i}. {step}")
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            style_run(run, size=10, color=TEXT)
    add_horizontal_rule(doc)

    add_heading(doc, "9. Ferramentas de desenvolvimento e qualidade")
    add_table(
        doc,
        ["Ferramenta", "Uso"],
        [
            ["Docker", "Supabase local (obrigatório para backend)"],
            ["Supabase CLI", "Migrations, functions, status"],
            ["Git", "Versionamento"],
            ["oxlint", "Linter no apps/web"],
            ["Vitest", "Testes dos schemas Zod"],
            ["React Native Testing Library", "Testes mobile (Fase 3)"],
            ["Cursor", "IDE com assistência de IA"],
        ],
    )
    add_horizontal_rule(doc)

    add_heading(doc, "10. Estado atual do desenvolvimento")
    add_status_table(
        doc,
        [
            ["Backend (migrations, RLS, storage)", "✅ Concluído"],
            ["Login + recuperação de senha", "✅ Concluído"],
            ["Gerenciador /admin", "✅ Concluído"],
            ["Plataforma aluno /app", "✅ Em evolução"],
            ["Leitor de livro paginado + layout imersivo", "✅ Testes em andamento"],
            ["App mobile (apps/mobile)", "⏳ Planejado"],
            ["Busca global, evolução do aluno, produções/obras", "⏳ Roadmap"],
        ],
    )
    add_horizontal_rule(doc)

    add_heading(doc, "11. Por que essas escolhas? (síntese para o TCC)")
    add_table(
        doc,
        ["Decisão", "Motivo"],
        [
            ["Supabase", "PostgreSQL + Auth + Storage + RLS sem montar backend do zero; adequado ao prazo do TCC"],
            ["React + Vite", "Ecossistema maduro, performance, deploy estático simples"],
            ["React Native + Expo", "App iOS/Android com paridade aluno; publicação futura via EAS"],
            ["TypeScript + Zod", "Segurança de tipos e validação compartilhada web/mobile"],
            ["Monorepo", "Web e mobile compartilham types e schemas"],
            ["TanStack Router/Query", "Rotas tipadas e cache sem Redux desnecessário"],
            ["Tailwind / NativeWind", "UI consistente entre web e app (RNF09)"],
        ],
    )
    add_horizontal_rule(doc)

    add_heading(doc, "12. Referências")
    refs = [
        ("Supabase Docs", "https://supabase.com/docs"),
        ("React", "https://react.dev"),
        ("React Native", "https://reactnative.dev"),
        ("Expo", "https://docs.expo.dev"),
        ("Vite", "https://vite.dev"),
        ("TanStack Router", "https://tanstack.com/router"),
        ("TanStack Query", "https://tanstack.com/query"),
        ("Tailwind CSS", "https://tailwindcss.com"),
        ("Zod", "https://zod.dev"),
    ]
    for name, url in refs:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{name}: ")
        style_run(r1, bold=True, size=10, color=BRAND)
        r2 = p.add_run(url)
        style_run(r2, size=10, color=ACCENT)

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(
        "Documento vivo — atualize conforme o apps/mobile for implementado e novas tecnologias entrarem no projeto."
    )
    style_run(r, italic=True, size=9, color=MUTED)

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Gerado: {OUTPUT}")


if __name__ == "__main__":
    main()
